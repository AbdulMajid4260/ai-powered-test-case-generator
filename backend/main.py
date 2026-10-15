
import logging
import os
from io import BytesIO
from uuid import uuid4

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError

from llm_service import generate_test_cases
from schemas import GenerateRequest, GenerateResponse, TestCase

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = FastAPI(
    title="AI-Powered Test Case Generator",
    version="1.1.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://127.0.0.1:3000",
        "http://localhost:5500",
        "http://127.0.0.1:5500",
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=False,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["Content-Type", "Authorization"],
)

# Temporary server-side storage for generated results.
generated_results: dict[str, list[TestCase]] = {}


@app.get("/health")
def health_check():
    return {"status": "healthy"}


@app.post(
    "/api/generate-test-cases",
    response_model=GenerateResponse,
)
def create_test_cases(request: GenerateRequest):
    requirement = request.requirement.strip()

    if not requirement:
        raise HTTPException(
            status_code=422,
            detail="Requirement must not be blank.",
        )

    try:
        test_cases = generate_test_cases(requirement)

        if not 5 <= len(test_cases) <= 10:
            raise ValueError("Generation must produce 5 to 10 test cases.")

        generation_id = str(uuid4())
        generated_results[generation_id] = test_cases

        return GenerateResponse(
            generation_id=generation_id,
            test_cases=test_cases,
        )

    except ValidationError:
        logger.warning("Generated test-case data failed validation.")
        raise HTTPException(
            status_code=502,
            detail="The AI returned invalid test-case data. Please try again.",
        ) from None

    except Exception as exc:
        # Never log the API key.
        logger.error(
            "Test-case generation failed (%s).",
            type(exc).__name__,
        )
        raise HTTPException(
            status_code=502,
            detail="Test-case generation failed. Please try again later.",
        ) from None


@app.get("/api/export/{generation_id}")
def export_test_cases(generation_id: str):
    test_cases = generated_results.get(generation_id)

    if test_cases is None:
        raise HTTPException(
            status_code=404,
            detail="Generated test cases were not found. Generate them again.",
        )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Test Cases"

    headers = [
        "Test Case ID",
        "Title",
        "Preconditions",
        "Test Steps",
        "Expected Result",
        "Priority",
        "Test Type",
    ]
    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

    for case in test_cases:
        worksheet.append([
            case.test_case_id,
            case.title,
            "\n".join(case.preconditions),
            "\n".join(case.test_steps),
            case.expected_result,
            case.priority,
            case.test_type,
        ])

    for column_cells in worksheet.columns:
        column_letter = column_cells[0].column_letter
        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )
        worksheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, 15),
            45,
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="test_cases_{generation_id}.xlsx"'
            )
        },
    )
